"""New functionality: dynamic TXT/MD parsing, live metadata, reports, hourly scheduler."""
import json
import sqlite3
import time
from pathlib import Path

import pytest
import yaml

from orion_db.schema_catalog import parse_schema
from orion_db.cli import create_demo_db
from orion_db.knowledge import redact_secrets


SAMPLE = '''# Database Schema Documentation
**Database Name:** AgentDB
**Engine:** Microsoft SQL Server
### ??? feedback
| Column | Type | Nullable | Default | Constraints | Description |
| :--- | :--- | :--- | :--- | :--- | :--- |
| `id` | `INT` | No | `*None*` | **PK** | - |
| `label` | `NVARCHAR(200)` | No | `*None*` | - | - |
---
**Database Name:** MFAssociate
**Engine:** Microsoft SQL Server
### ??? Agencies
| Column | Type | Nullable | Default | Constraints | Description |
| :--- | :--- | :--- | :--- | :--- | :--- |
| `AgencyId` | `INT` | No | `*None*` | **PK** | - |
'''


def test_parse_multiple_databases_and_columns():
    result = parse_schema(SAMPLE, 'schema.md')
    assert set(result) == {'agentdb', 'mfassociate'}
    assert list(result['agentdb']['tables']) == ['feedback']
    assert result['agentdb']['tables']['feedback']['columns'][1]['name'] == 'label'
    assert result['mfassociate']['tables']['agencies']['columns'][0]['constraints'] == 'PK'


def test_incremental_md_and_deleted_source(identity):
    app, auth, root = identity
    source = root / 'knowledge' / 'schemas.md'
    source.write_text(SAMPLE, encoding='utf-8')
    first = app.discovery.refresh(documents=True, schemas=False)
    assert first['status'] == 'completed'
    catalog = app.process({**auth, 'type': 'schema_catalog'})['databases']
    assert {'AgentDB', 'MFAssociate'} <= {x['name'] for x in catalog}
    assert app.knowledge.sync_once() == 'unchanged'
    source.write_text(SAMPLE.replace('### ??? feedback', '### ??? interactions'), encoding='utf-8')
    app.discovery.refresh(documents=True, schemas=False)
    agent = next(x for x in app.discovery.catalog.entries() if x['name'] == 'AgentDB')
    assert [t['name'] for t in agent['tables']] == ['interactions']
    source.unlink()
    app.discovery.refresh(documents=True, schemas=False)
    assert 'AgentDB' not in {x['name'] for x in app.discovery.catalog.entries()}


def test_secret_redaction_in_index(identity):
    app, auth, root = identity
    source = root / 'knowledge' / 'credentials.md'
    source.write_text('Database: safe_db\nPassword: HighlySecret123\nPWD=AnotherSecret\n', encoding='utf-8')
    app.discovery.refresh(documents=True, schemas=False)
    content = (root / 'index.json').read_text()
    assert 'HighlySecret123' not in content
    assert 'AnotherSecret' not in content
    assert 'HighlySecret123' not in json.dumps(app.discovery.catalog.entries())


def test_discovery_status_and_admin_refresh(identity):
    app, auth, root = identity
    result = app.process({**auth, 'type': 'refresh_discovery'})
    assert result['type'] == 'discovery_status'
    assert result['status'] == 'completed'
    assert result['catalog']['documented_databases'] >= 1
    status = app.process({**auth, 'type': 'discovery_status'})
    assert status['knowledge']['sync_interval_seconds'] == 3600


def test_verified_sqlite_schema_and_excel(identity):
    app, auth, root = identity
    path = create_demo_db(root / 'inventory.sqlite')
    data = yaml.safe_load(app.cfg.path.read_text())
    data['database']['profiles'] = [{'alias': 'inventory', 'backend': 'sqlite', 'database': str(path)}]
    app.cfg.path.write_text(yaml.safe_dump(data, sort_keys=False))
    app.discovery.refresh(documents=False, schemas=True)
    catalog = app.discovery.catalog.entries()
    inventory = next(x for x in catalog if x['name'] == str(path))
    assert inventory['verified']
    assert {t['name'] for t in inventory['tables']} == {'assets', 'locations'}
    assert len(next(t for t in inventory['tables'] if t['name'] == 'assets')['columns']) == 4
    result = app.process({**auth, 'type': 'db_connect_saved', 'alias': 'inventory'})
    assert result['type'] == 'db_connected'
    report = app.process({**auth, 'type': 'report_generate', 'schema': 'main', 'table': 'assets'})
    assert report['type'] == 'report_ready'
    assert report['row_count'] == 3
    assert 'SELECT' in report['sql']
    from openpyxl import load_workbook
    book = load_workbook(report['path'], read_only=True)
    rows = list(book['assets'].values)
    assert rows[0] == ('asset_id', 'asset_name', 'department', 'purchase_year')
    assert rows[1][1] == 'Laptop A'
    assert book['Report Info']['B5'].value == 3


def test_export_rejects_unknown_table(identity):
    app, auth, root = identity
    path = create_demo_db(root / 'test.sqlite')
    app.process({**auth, 'type': 'db_connect', 'backend': 'sqlite', 'database': str(path), 'confirmed': True})
    with pytest.raises((ValueError, PermissionError)):
        app.process({**auth, 'type': 'report_generate', 'schema': 'main', 'table': 'assets; DROP TABLE assets'})


def test_document_report_requires_actual_rows(identity):
    app, auth, root = identity
    (root / 'knowledge' / 'schema.md').write_text(SAMPLE, encoding='utf-8')
    (root / 'knowledge' / 'actual.md').write_text('| Department | Amount |\n|---|---|\n| Sales | 250 |\n| IT | 120 |\n', encoding='utf-8')
    app.discovery.refresh(documents=True, schemas=False)
    with pytest.raises(ValueError, match='No actual tabular data'):
        app.process({**auth, 'type': 'report_generate_document', 'source': 'schema.md'})
    report = app.process({**auth, 'type': 'report_generate_document', 'source': 'actual.md'})
    assert report['row_count'] == 2
    from openpyxl import load_workbook
    book = load_workbook(report['path'], read_only=True)
    assert list(book.active.values)[1] == ('Sales', '250')


def test_scheduler_starts_and_stops(identity):
    app, auth, root = identity
    app.discovery.start()
    try:
        for _ in range(30):
            if app.discovery.status()['scheduler_running']:
                break
            time.sleep(.01)
        assert app.discovery.status()['scheduler_running']
    finally:
        app.discovery.stop()
    assert not app.discovery.status()['scheduler_running']


def test_general_question_not_misclassified_as_database_list():
    from orion_db.server import wants_catalog
    from orion_db.routing import route_question
    assert not wants_catalog('what is database')
    assert wants_catalog('list of database we have')
    assert route_question('what is database', {'ai': {'mode': 'hybrid'}}) == 'general'


def test_connect_to_documented_database_requests_details(identity):
    app, auth, root = identity
    (root / 'knowledge' / 'agent.md').write_text('**Database Name:** AgentDB\n### ??? feedback\n', encoding='utf-8')
    app.discovery.refresh(documents=True, schemas=False)
    tid = app.process({**auth, 'type': 'new_chat'})['thread_id']
    result = app.process({**auth, 'type': 'chat', 'thread_id': tid, 'message': 'connect to AgentDB'})
    assert result['type'] == 'db_connect_required'
    assert not app.live.status(auth['token'])['connected']
