"""Incremental knowledge sync with source-aware retrieval and file date metadata."""
import hashlib
import json
import os
import re
import threading
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

from .config import resolve_path
from .embeddings import SemanticIndex

STOP = set('a an the is are of to in on for and or with what which where how do does can i we you me our please show tell about have databases db data information give list details'.split())
WORDS = re.compile(r'[a-z0-9_]+', re.I)
# Hide actual credential values in the searchable index and schema cache.
SECRET_LINE = re.compile(r'(?im)^(\s*(?:\*\*)?\s*(?:db[_ -]?password|database[_ -]?password|password|pwd|passphrase|secret|api[_ -]?key|access[_ -]?token|connection[_ -]?string)(?:\*\*)?\s*[:=]\s*).+$')
SECRET_INLINE = re.compile(r'(?i)\b(PWD|Password|Secret|Token)\s*=\s*([^;\s]+)')
RESUME_NAME = re.compile(r'(?i)(?:^|[/_\\ .-])(?:resumes?|cvs?|candidates?|curriculum[ _-]?vitae)(?:$|[/_\\ .-])')
RESUME_HEADER = re.compile(r'(?im)^\s*(?:resume|curriculum vitae|professional (?:profile|summary))\s*$')
RESUME_HINTS = re.compile(r'(?i)\b(work experience|professional experience|employment history|education|skills|career objective|projects)\b')


def terms(text):
    return [w for w in WORDS.findall(str(text).casefold()) if len(w) > 1 and w not in STOP]


def redact_secrets(text):
    text = SECRET_LINE.sub(lambda m: m.group(1) + '[REDACTED]', text)
    return SECRET_INLINE.sub(lambda m: m.group(1) + '=[REDACTED]', text)


class Knowledge:
    def __init__(self, cfg, audit):
        self.cfg, self.audit = cfg, audit
        self._lock = threading.RLock()
        self._index = {'files': [], 'chunks': [], 'snapshot': [], 'document_types': {}}
        self.last_status = 'not_started'
        self.last_error = ''
        self.last_sync = ''
        self._semantic = SemanticIndex(cfg)
        self.classifier = None  # Attached by App after the local Qwen3 model is constructed.
        self._load_previous()

    def _settings(self):
        return self.cfg.get()['knowledge']

    def _folder(self):
        return resolve_path(self._settings()['folder'])

    def _index_path(self):
        return resolve_path(self._settings()['index_file'])

    def _load_previous(self):
        path = self._index_path()
        if path.is_file():
            try:
                value = json.loads(path.read_text(encoding='utf-8'))
                if isinstance(value.get('chunks'), list):
                    self._index = value
                    self.last_status = 'loaded_previous_index'
                    self.last_sync = value.get('last_sync', '')
            except (OSError, ValueError, TypeError) as exc:
                self.audit.error('knowledge.load', exc)

    def _snapshot(self):
        root = self._folder()
        if not root.is_dir():
            raise FileNotFoundError('Knowledge directory does not exist: ' + str(root))
        settings = self._settings()
        allowed = {x.lower() for x in settings.get('supported_extensions', ['.txt', '.md'])}
        result = []
        for path in sorted(root.rglob('*')):
            if not path.is_file() or path.is_symlink() or path.suffix.lower() not in allowed or path.name.startswith('~$'):
                continue
            size = path.stat().st_size
            if size > int(settings.get('max_file_bytes', 500000)):
                raise ValueError('Knowledge document too large: ' + str(path))
            if len(result) >= int(settings.get('max_files', 1000)):
                raise ValueError('Knowledge file limit exceeded')
            h = hashlib.sha256(path.read_bytes()).hexdigest()
            stat = path.stat()
            # Windows: file creation timestamp; POSIX: birthtime when available.
            # Some filesystems do not expose creation time. Fall back to mtime,
            # and explicitly report the weaker timestamp basis to users.
            birth = getattr(stat, 'st_birthtime', None)
            if birth is not None:
                created, basis = birth, 'filesystem_creation'
            elif os.name == 'nt':
                created, basis = stat.st_ctime, 'windows_creation'
            else:
                created, basis = stat.st_mtime, 'modified_time_fallback'
            result.append({
                'source': path.relative_to(root).as_posix(), 'sha256': h,
                'created_at': datetime.fromtimestamp(created, timezone.utc).isoformat(),
                'modified_at': datetime.fromtimestamp(stat.st_mtime, timezone.utc).isoformat(),
                'created_time_basis': basis,
            })
        return result

    def _read_document(self, path):
        """Extract text from supported local formats. No OCR or external upload."""
        suffix = path.suffix.lower()
        if suffix in ('.txt', '.md'):
            return path.read_text(encoding='utf-8-sig', errors='replace')
        if suffix == '.pdf':
            try:
                from pypdf import PdfReader
            except ImportError as exc:
                raise RuntimeError('PDF support requires: pip install pypdf') from exc
            reader = PdfReader(str(path))
            max_pages = int(self._settings().get('max_pdf_pages', 40))
            return '\n'.join((page.extract_text() or '') for page in list(reader.pages)[:max_pages])
        if suffix == '.docx':
            try:
                from docx import Document
            except ImportError as exc:
                raise RuntimeError('DOCX support requires: pip install python-docx') from exc
            document = Document(str(path))
            lines = [p.text for p in document.paragraphs]
            for table in document.tables:
                lines.extend(' | '.join(c.text for c in row.cells) for row in table.rows)
            return '\n'.join(lines)
        raise ValueError('Unsupported knowledge format: ' + suffix)

    @staticmethod
    def _document_type(source, text):
        # Filename/folder hints work for scanned PDFs too; content hints are
        # useful when resumes are dropped directly into the knowledge root.
        if RESUME_NAME.search(source) or RESUME_HEADER.search(text[:3000]):
            return 'resume'
        if len(set(x.casefold() for x in RESUME_HINTS.findall(text[:12000]))) >= 2:
            return 'resume'
        return 'document'

    def resume_inventory(self, include_names=False, allowed_sources=None):
        """Deterministic counts from synchronized metadata, never Qwen3."""
        with self._lock:
            types = self._index.get('document_types', {})
            names = sorted((s for s, kind in types.items() if kind == 'resume' and
                            (allowed_sources is None or s in allowed_sources)), key=str.casefold)
            result = {'count': len(names), 'last_sync': self.last_sync,
                      'folder': str(self._folder()), 'total_indexed_files': len(self._index['files']) if allowed_sources is None else len(allowed_sources)}
            if include_names:
                result['files'] = names
            return result

    def recent_inventory(self, document_type='all', activity='created', date=None, allowed_sources=None):
        """Query indexed filesystem timestamps; no LLM-generated dates or counts.

        Dates use the SERVER's local timezone. A file's creation timestamp
        is not the date a recruiter received a resume unless that is when it
        was saved on this server. On filesystems without birthtime, creation
        is approximated using the last-modified timestamp.
        """
        if document_type not in ('all', 'resume', 'document'):
            raise ValueError('Invalid document type')
        if activity not in ('created', 'modified', 'all'):
            raise ValueError('Invalid document activity')
        target_date = date or datetime.now().astimezone().date()
        with self._lock:
            document_types = self._index.get('document_types', {})
            matched = []
            for entry in self._index.get('snapshot', []):
                source = entry['source']
                if allowed_sources is not None and source not in allowed_sources:
                    continue
                kind = document_types.get(source, 'document')
                if document_type != 'all' and kind != document_type:
                    continue
                created_at = entry.get('created_at')
                modified_at = entry.get('modified_at')
                # Old indexes have no timestamps: require a fresh sync.
                if not created_at or not modified_at:
                    continue
                created_date = datetime.fromisoformat(created_at).astimezone().date()
                modified_date = datetime.fromisoformat(modified_at).astimezone().date()
                if not ((activity in ('created', 'all') and created_date == target_date) or
                        (activity in ('modified', 'all') and modified_date == target_date)):
                    continue
                matched.append({
                    'source': source, 'document_type': kind,
                    'created_at': created_at, 'modified_at': modified_at,
                    'created_time_basis': entry.get('created_time_basis', 'unknown'),
                })
            matched.sort(key=lambda item: item['source'].casefold())
            return {
                'date': target_date.isoformat(), 'activity': activity,
                'document_type': document_type, 'count': len(matched),
                'files': matched, 'last_sync': self.last_sync,
                'timezone': str(datetime.now().astimezone().tzinfo),
                'timestamp_fallback_used': any(x['created_time_basis'] == 'modified_time_fallback'
                                               for x in matched),
            }

    def sync_once(self):
        with self._lock:
            try:
                before = self._snapshot()
                if before == self._index.get('snapshot', []) and self._index.get('version', 0) >= 5:
                    # Pending classifications can be retried after Qwen3 is installed.
                    self._sync_classifications()
                    self.last_status, self.last_error = 'unchanged', ''
                    if not self.last_sync:
                        self.last_sync = datetime.now(timezone.utc).isoformat()
                    return 'unchanged'
                settings = self._settings()
                root = self._folder()
                chunk_size = max(250, int(settings.get('chunk_chars', 1200)))
                overlap = min(chunk_size - 1, max(0, int(settings.get('chunk_overlap', 150))))
                previous_hashes = {x['source']: x['sha256'] for x in self._index.get('snapshot', [])}
                previous_chunks = {}
                for chunk in self._index.get('chunks', []):
                    previous_chunks.setdefault(chunk['source'], []).append(chunk)
                chunks = []
                document_types = {}
                previous_types = self._index.get('document_types', {})
                changed = 0
                for file in before:
                    source = file['source']
                    if (self._index.get('version', 0) >= 4 and previous_hashes.get(source) == file['sha256']
                            and source in previous_types):
                        document_types[source] = previous_types[source]
                        chunks.extend(previous_chunks.get(source, []))
                        continue
                    changed += 1
                    text = self._read_document(root / source)
                    document_types[source] = self._document_type(source, text)
                    if settings.get('redact_connection_secrets_before_indexing', True):
                        text = redact_secrets(text)
                    if not text.strip():
                        continue
                    pos = 0
                    while pos < len(text):
                        segment = text[pos:pos + chunk_size]
                        chunks.append({'source': source, 'text': segment})
                        if pos + chunk_size >= len(text):
                            break
                        pos += chunk_size - overlap
                if before != self._snapshot():
                    raise RuntimeError('Documents changed during sync; retry')
                now = datetime.now(timezone.utc).isoformat()
                index = {'version': 5, 'snapshot': before, 'files': [f['source'] for f in before],
                         'document_types': document_types, 'chunks': chunks, 'last_sync': now}
                target = self._index_path()
                target.parent.mkdir(parents=True, exist_ok=True)
                temp = target.with_suffix(target.suffix + '.tmp')
                try:
                    temp.write_text(json.dumps(index, ensure_ascii=False), encoding='utf-8')
                    os.replace(temp, target)
                finally:
                    temp.unlink(missing_ok=True)
                self._index = index
                self._sync_classifications()
                self.last_status, self.last_error, self.last_sync = 'updated', '', now
                self.audit.event('knowledge_sync', files=len(before), changed_files=changed,
                                 chunks=len(chunks), status='updated')
                return 'updated'
            except Exception as exc:
                self.last_status, self.last_error = 'error', str(exc)
                self.audit.error('knowledge.sync', exc)
                self.audit.event('knowledge_sync', status='error', error=str(exc))
                return 'error'

    def _sync_classifications(self):
        """Classification is advisory and must not prevent knowledge sync."""
        if self.classifier is None:
            return
        try:
            self.classifier.sync(self._index.get('snapshot', []), self._index.get('chunks', []))
        except Exception as exc:
            self.audit.event('document_classification_sync_failed', error_type=type(exc).__name__)

    def read_source_for_role(self, source, role):
        """Read only one authorized live document, never the whole catalog.

        Used for document Excel exports. Enforce the SHA256-bound grant both
        before and after the read, so a file edit invalidates the result.
        """
        if getattr(self, 'access_control', None) is None:
            raise PermissionError('Document access enforcement is unavailable')
        self.access_control.require(role, self, source)
        root = self._folder().resolve()
        original = root / source
        path = original.resolve()
        if original.is_symlink() or not path.is_relative_to(root) or not path.is_file():
            raise PermissionError('Document source is unavailable')
        text = redact_secrets(self._read_document(path))
        self.access_control.require(role, self, source)
        return text

    def source_texts(self):
        """Return redacted documents for schema parsing (not live records)."""
        with self._lock:
            result = {}
            root = self._folder()
            for source in self._index.get('files', []):
                path = root / source
                try:
                    if path.is_symlink() or not path.resolve().is_relative_to(root.resolve()):
                        continue
                    text = self._read_document(path)
                    result[source] = redact_secrets(text)
                except OSError:
                    continue
            return result

    def status(self):
        with self._lock:
            classification = None
            if self.classifier is not None:
                try:
                    classification = self.classifier.status()
                except (OSError, ValueError, TypeError):
                    classification = {'enabled': True, 'status': 'unavailable'}
            return {'classification': classification, 'status': self.last_status, 'error': self.last_error, 'last_sync': self.last_sync,
                    'folder': str(self._folder()), 'files': len(self._index['files']),
                    'chunks': len(self._index['chunks']),
                    'resume_files': self.resume_inventory()['count'],
                    'background_monitoring': bool(int(self._settings().get('sync_interval_seconds', 0)) > 0),
                    'sync_interval_seconds': int(self._settings().get('sync_interval_seconds', 0))}

    def catalog(self):
        """Documented database names only; does not imply a verified connection."""
        found = {}
        with self._lock:
            for source, text in self.source_texts().items():
                for match in re.finditer(r'(?im)^\s*(?:\*\*)?(?:database name|database)(?:\*\*)?\s*:\s*(?:\*\*)?\s*([\w.\-]{2,100})', text):
                    name = match.group(1)
                    found.setdefault(name.casefold(), {'name': name, 'source': source})
        return sorted(found.values(), key=lambda x: x['name'].casefold())

    def search(self, question, limit=None, include_resumes=True, allowed_sources=None):
        q = Counter(terms(question))
        if not q:
            return []
        limit = max(1, min(10, int(limit or self._settings().get('max_results', 4))))
        scored = []
        with self._lock:
            for chunk in self._index['chunks']:
                if allowed_sources is not None and chunk['source'] not in allowed_sources:
                    continue
                if not include_resumes and self._index.get('document_types', {}).get(chunk['source']) == 'resume':
                    continue
                words = Counter(terms(chunk['text']))
                matched = sum(1 for t in q if words[t])
                if not matched:
                    continue
                score = sum(min(4, words[t]) * (1.0 + min(2, q[t]) * .25) for t in q)
                score += matched * 2
                score += sum(2.5 for t in q if t in chunk['source'].casefold())
                if len(q) > 1 and matched == len(q):
                    score += 4
                scored.append((score, chunk))
        scored.sort(key=lambda x: x[0], reverse=True)
        rag = self.cfg.get().get('rag', {})
        if rag.get('enabled', False) and rag.get('retrieval') == 'hybrid':
            # Combine lexical and local semantic rankings using reciprocal rank fusion.
            # If the optional embedding model is unavailable, retain lexical RAG.
            with self._lock:
                chunks = list(self._index['chunks'])
                types = dict(self._index.get('document_types', {}))
            semantic_sources = {c['source'] for c in chunks if (allowed_sources is None or c['source'] in allowed_sources) and
                               (include_resumes or types.get(c['source']) != 'resume')}
            try:
                semantic = self._semantic.search(question, chunks, semantic_sources, limit * 4)
                ranks = {}
                chunks_by_id = {}
                for ranking, items in ((0, scored), (1, semantic)):
                    for rank, (_, chunk) in enumerate(items[:max(20, limit * 4)], 1):
                        key = (chunk['source'], chunk['text'])
                        chunks_by_id[key] = chunk
                        ranks.setdefault(key, [0.0, 0.0])[ranking] = 1.0 / (60 + rank)
                weight = max(0, min(1, float(rag.get('semantic_weight', 0.6))))
                combined = sorted(((weight * scores[1] + (1-weight) * scores[0], chunks_by_id[key])
                                   for key, scores in ranks.items()), reverse=True, key=lambda item: item[0])
                # Defense in depth: even if a semantic backend returns a
                # foreign chunk, it cannot enter the retrieval result.
                allowed_pairs = {(c['source'], c['text']) for c in chunks
                                 if c['source'] in semantic_sources}
                return [{'source': chunk['source'], 'score': round(score, 5), 'text': chunk['text']}
                        for score, chunk in combined if
                        (chunk['source'], chunk['text']) in allowed_pairs][:limit]
            except (ImportError, RuntimeError, OSError, ValueError, FileNotFoundError) as exc:
                self.audit.event('semantic_fallback_lexical', error_type=type(exc).__name__)
        return [{'source': chunk['source'], 'score': round(score, 2), 'text': chunk['text']}
                for score, chunk in scored[:limit]]

    def document_evidence(self, question, allowed_sources, document_type='all',
                          max_sources=12, max_lines_per_source=6):
        """Collect source-level excerpts from approved indexed documents.

        Unlike chunk top-k, this can cover multiple documents. It does not
        infer missing skills, names, or facts. Caller must enforce the ACL.
        """
        if allowed_sources is None:
            raise ValueError('Explicit authorized source set is required')
        if document_type not in ('all', 'resume', 'document'):
            raise ValueError('Unsupported document type')
        query_terms = set(terms(question))
        if not query_terms:
            return []
        with self._lock:
            types = dict(self._index.get('document_types', {}))
            grouped = {}
            for chunk in self._index.get('chunks', []):
                source = chunk.get('source')
                if source not in allowed_sources or (document_type != 'all' and
                    types.get(source, 'document') != document_type):
                    continue
                grouped.setdefault(source, []).append(chunk.get('text', ''))
        evidence = []
        for source, chunks in grouped.items():
            # Merge overlapping chunks so a line cannot be counted twice.
            lines = []
            seen = set()
            for chunk in chunks:
                for raw in chunk.splitlines():
                    line = raw.strip(' \t-•')
                    if not 3 <= len(line) <= 350 or line in seen:
                        continue
                    seen.add(line)
                    line_terms = set(terms(line))
                    overlap = len(query_terms & line_terms)
                    # Skill/stack questions benefit from explicit skill sections
                    # even if the user did not name the exact technology.
                    skill_request = bool(query_terms & {'skill', 'skills', 'skillset',
                        'techstack', 'technologies', 'technology', 'stack'})
                    skill_line = bool(re.search(r'(?i)\b(skills?|technologies|tech stack|'
                        r'frameworks?|programming languages?|tools?)\s*:', line))
                    score = overlap * 3 + (2 if skill_request and skill_line else 0)
                    if score:
                        lines.append((score, line))
            lines.sort(key=lambda item: item[0], reverse=True)
            if lines:
                evidence.append({'source': source, 'score': lines[0][0],
                    'excerpts': [line for _, line in lines[:max_lines_per_source]]})
        evidence.sort(key=lambda item: (-item['score'], item['source'].casefold()))
        return evidence[:max_sources]

    def approved_answer(self, question, allowed_sources=None):
        """Use exact, administrator-approved corrections, not raw chat history."""
        normalize = lambda x: re.sub(r'[^\w]+', ' ', x.casefold()).strip()
        wanted = normalize(question)
        with self._lock:
            sources = [name for name in self._index.get('files', [])
                       if name.startswith('Approved_Feedback/') and name.endswith('.md') and
                       (allowed_sources is None or name in allowed_sources)]
            for source in sources:
                path = self._folder() / source
                try:
                    if path.is_symlink() or not path.resolve().is_relative_to(self._folder().resolve()):
                        continue
                    text = redact_secrets(path.read_text(encoding='utf-8'))
                except OSError:
                    continue
                found = re.search(r'(?ms)^Question:\s*(.*?)\n\s*Answer:\s*(.*?)\s*$', text)
                if found and normalize(found.group(1)) == wanted:
                    return found.group(2).strip(), source
        return None

    def answer_from_docs(self, question, hits):
        """Short extractive answer: never copy whole tables or fabricate records."""
        if not hits:
            return 'I could not find that information in the synchronized TXT/MD documents.'
        query = set(terms(question))
        candidates = []
        for hit in hits:
            for raw in hit['text'].splitlines():
                line = raw.strip(' \t-•')
                # Table metadata is useful for discovery but not a prose answer.
                if (not 12 <= len(line) <= 300 or line.startswith('|') or
                    line.startswith('###') or line.startswith('---') or
                    line.startswith('```') or 'password' in line.casefold()):
                    continue
                overlap = len(query.intersection(terms(line)))
                if overlap:
                    candidates.append((overlap, hit['score'], line, hit['source']))
        if not candidates:
            return ('The synchronized documents mention this topic, but do not give a clear answer. '
                    'Source: ' + hits[0]['source'])
        candidates.sort(key=lambda x: (x[0], x[1]), reverse=True)
        _, _, line, source = candidates[0]
        # Do not let a retrieved file issue instructions to the assistant.
        line = re.sub(r'(?i)\b(ignore previous instructions|system prompt|assistant role)\b', '[untrusted text]', line)
        return line + ' (Source: ' + source + ')'
