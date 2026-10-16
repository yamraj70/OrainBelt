"""Generate bounded Excel exports from verified live tables (not schema docs)."""
import re
import uuid
from datetime import datetime, timezone
from pathlib import Path

from .config import resolve_path
from .database import ident

SENSITIVE = re.compile(r'(?i)(password|passwd|secret|api[_-]?key|access[_-]?token|passwordhash|aadhaar|ssn|private[_-]?key)')


def generate_table_report(cfg, audit, live, token, schema, table):
    """Read-only SQL, checked identifiers and permissions, export actual rows to XLSX."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise RuntimeError('Excel reports require: pip install openpyxl') from exc
    session = live._session(token)
    backend = session['info']['backend']
    settings = cfg.get().get('reports', {})
    maximum = max(1, min(100000, int(settings.get('query', {}).get('max_export_rows', 50000))))
    batch = max(1, min(5000, int(settings.get('query', {}).get('fetch_batch_size', 1000))))
    with session['lock']:
        if {'schema': schema, 'table': table} not in live.tables(token):
            raise PermissionError('Table is not present in verified live schema')
        target = ident(table, backend)
        if backend != 'sqlite':
            target = ident(schema, backend) + '.' + target
        cursor = session['conn'].cursor()
        # Inspect only column names, then select non-sensitive columns explicitly.
        cursor.execute(f'SELECT TOP 0 * FROM {target}' if backend == 'sqlserver'
                       else f'SELECT * FROM {target} LIMIT 0')
        names = [col[0] for col in cursor.description]
        visible = [i for i, name in enumerate(names) if not SENSITIVE.search(name)]
        if not visible:
            raise PermissionError('No exportable columns: all columns are restricted')
        select_columns = ', '.join(ident(names[i], backend) for i in visible)
        sql = (f'SELECT TOP {maximum + 1} {select_columns} FROM {target}' if backend == 'sqlserver'
               else f'SELECT {select_columns} FROM {target} LIMIT {maximum + 1}')
        cursor.execute(sql)
        wb = Workbook(write_only=False)
        sheet = wb.active
        sheet.title = re.sub(r'[\\/*?:\[\]]', '_', table)[:31] or 'Report'
        sheet.append([names[i] for i in visible])
        count = 0
        truncated = False
        while True:
            rows = cursor.fetchmany(batch)
            if not rows:
                break
            for row in rows:
                if count >= maximum:
                    truncated = True
                    break
                values = []
                for i in visible:
                    cell = row[visible.index(i)]
                    if isinstance(cell, (bytes, bytearray)):
                        cell = '[binary]'
                    if isinstance(cell, str) and cell.lstrip().startswith(('=', '+', '-', '@')):
                        cell = "'" + cell  # avoid spreadsheet formula injection
                    if cell is not None and not isinstance(cell, (int, float, str, bool, datetime)):
                        cell = str(cell)[:32000]
                    values.append(cell)
                sheet.append(values)
                count += 1
            if truncated:
                break
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='20304A')
        cell.alignment = Alignment(wrap_text=True)
    for index, name in enumerate([names[i] for i in visible], start=1):
        sheet.column_dimensions[get_column_letter(index)].width = min(48, max(14, len(name) + 3))
    if settings.get('include_metadata_sheet', True):
        meta = wb.create_sheet('Report Info')
        for key, value in [('Database', session['info']['database']), ('Schema', schema),
                           ('Table', table), ('Generated UTC', datetime.now(timezone.utc).isoformat()),
                           ('Rows', count), ('Truncated', truncated), ('SQL', sql)]:
            meta.append([key, value])
        meta.column_dimensions['A'].width = 22
        meta.column_dimensions['B'].width = 70
    output_dir = resolve_path(settings.get('output_directory', 'reports/generated'))
    output_dir.mkdir(parents=True, exist_ok=True)
    filename = re.sub(r'[^A-Za-z0-9_-]', '_', session['info']['database'] + '_' + table)[:90]
    output = output_dir / (filename + '_' + datetime.now(timezone.utc).strftime('%Y%m%d_%H%M%S') + '_' + uuid.uuid4().hex[:6] + '.xlsx')
    wb.save(output)
    audit.event('report_generated', username=session['user'], database=session['info']['database'],
                schema=schema, table=table, row_count=count, truncated=truncated)
    return {'type': 'report_ready', 'database': session['info']['database'], 'schema': schema,
            'table': table, 'sql': sql, 'row_count': count, 'truncated': truncated,
            'excluded_columns': len(names) - len(visible), 'path': str(output)}


def generate_document_report(cfg, audit, knowledge, source, role=None):
    """Export actual tabular rows in a synchronized TXT/MD file, not schema definitions."""
    import csv
    from io import StringIO
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise RuntimeError('Excel reports require: pip install openpyxl') from exc
    if getattr(knowledge, 'access_control', None) is not None:
        knowledge.access_control.require(role, knowledge, source)
    if source not in knowledge._index.get('files', []):
        raise PermissionError('Document is not in the approved synchronized knowledge index')
    root = knowledge._folder().resolve()
    path = (root / source).resolve()
    if not path.is_relative_to(root) or not path.is_file():
        raise PermissionError('Invalid document source path')
    text = knowledge.read_source_for_role(source, role)
    # Markdown table parser; ignore heading/metadata and schema-definition tables.
    lines = text.splitlines()
    tables = []
    for i in range(len(lines) - 2):
        if not lines[i].lstrip().startswith('|'):
            continue
        if not re.fullmatch(r'[\s|:\-]+', lines[i + 1]):
            continue
        headers = [v.strip().strip('`') for v in lines[i].strip().strip('|').split('|')]
        if not all(headers) or len(headers) < 2:
            continue
        rows = []
        for raw in lines[i + 2:]:
            if not raw.lstrip().startswith('|'):
                break
            values = [v.strip().strip('`') for v in raw.strip().strip('|').split('|')]
            if len(values) == len(headers):
                rows.append(values)
        if rows:
            tables.append((headers, rows))
    if not tables and path.suffix.lower() == '.txt':
        try:
            dialect = csv.Sniffer().sniff(text[:8192], delimiters=',\t|')
            records = list(csv.reader(StringIO(text), dialect))
            if len(records) > 1 and len(records[0]) >= 2:
                tables.append((records[0], [row for row in records[1:] if len(row) == len(records[0])]))
        except csv.Error:
            pass
    tables = [(headers, rows) for headers, rows in tables
              if not ({'column', 'type', 'nullable'} <= {h.casefold() for h in headers})]
    if not tables:
        raise ValueError('No actual tabular data found. This may be schema documentation; connect to the live database to export records.')
    limit = max(1, min(100000, int(cfg.get().get('reports', {}).get('query', {}).get('max_export_rows', 50000))))
    wb = Workbook()
    wb.remove(wb.active)
    count = 0
    for n, (headers, rows) in enumerate(tables[:20], 1):
        visible = [i for i, h in enumerate(headers) if not SENSITIVE.search(h)]
        if not visible:
            continue
        ws = wb.create_sheet('Data ' + str(n))
        ws.append([headers[i] for i in visible])
        for row in rows[:limit]:
            values = []
            for i in visible:
                value = row[i]
                if value.lstrip().startswith(('=', '+', '-', '@')):
                    value = "'" + value
                values.append(value[:32000])
            ws.append(values)
            count += 1
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='20304A')
        for i, col in enumerate(visible, 1):
            ws.column_dimensions[get_column_letter(i)].width = min(48, max(14, len(headers[col]) + 3))
    if not wb.worksheets:
        raise PermissionError('All discovered columns are restricted from export')
    out_dir = resolve_path(cfg.get().get('reports', {}).get('output_directory', 'reports/generated'))
    out_dir.mkdir(parents=True, exist_ok=True)
    filename = re.sub(r'[^A-Za-z0-9_-]', '_', path.stem)[:80]
    out = out_dir / (filename + '_' + uuid.uuid4().hex[:8] + '.xlsx')
    wb.save(out)
    audit.event('document_report_generated', source=source, row_count=count)
    return {'type': 'report_ready', 'source': source, 'row_count': count, 'path': str(out),
            'sql': '', 'message': 'Exported actual document table rows; no live SQL executed.'}
